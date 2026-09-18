#!/usr/bin/env python3
"""
Minimal markdown-to-docx converter for use when pandoc isn't installed.

Handles: #..###### headers, **bold**, `code` spans, - bullets, 1. numbered
lists, and bare *italic* lines (e.g. a closing colophon). Not a full
CommonMark implementation — covers the common shape of synthesis/report
markdown (headers + prose + lists + emphasis), which is the vast majority
of AI-agent-generated report content.

Usage:
    python3 md_to_docx.py --src input.md --out output.docx

After running, verify structurally (see document-layout-design SKILL.md,
"Converting an Existing Markdown Document to DOCX" section) — compare the
heading list and word count against the source before treating the
conversion as done.
"""
import argparse
import re
from docx import Document
from docx.shared import Pt

BOLD_OR_CODE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_runs(paragraph, text):
    """Add text to paragraph, handling **bold** and `code` inline markup."""
    parts = BOLD_OR_CODE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def convert(src_path, out_path):
    with open(src_path, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    i = 0
    n = len(lines)
    while i < n:
        stripped = lines[i].strip()

        if stripped == "---":
            i += 1
            continue
        if stripped == "":
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            h = doc.add_heading(level=level)
            add_runs(h, m.group(2))
            i += 1
            continue

        m = re.match(r"^-\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(2))
            i += 1
            continue

        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(stripped[1:-1])
            run.italic = True
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--src", required=True, help="Source markdown file path")
    parser.add_argument("--out", required=True, help="Output .docx file path")
    args = parser.parse_args()
    result = convert(args.src, args.out)
    print(f"Wrote {result}")
